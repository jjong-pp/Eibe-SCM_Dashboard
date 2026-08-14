"""
엑셀 입출력 테스트.

구 excel_parser.py 는 `except Exception: continue` 로 잘못된 행을 조용히
버렸고, 날짜와 금액을 문자열로 뭉갰다. 여기서 고정하는 것은:
  - 값 변환이 실제 엑셀이 주는 타입(datetime · float · 일련번호)을 받는가
  - 금액이 float 을 거치지 않는가 (P-02 가 입력 단계에서 되풀이되지 않게)
  - 잘못된 행이 조용히 사라지지 않고 위치와 함께 보고되는가
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.enums import InboundStatus
from app.models.metrics import WeeklyMetric
from app.models.sales import Promotion, SalesOrder
from app.models.scm import Inbound, InventorySnapshot
from app.services import excel as E
from tests.factories import (
    make_alias,
    make_brand,
    make_channel,
    make_product,
    make_warehouse,
)


def build_sheet(kind: E.Kind, rows: list[list]) -> bytes:
    """헤더가 있는 실제 xlsx 를 만든다."""
    template = E.TEMPLATES[kind]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = template.sheet_name
    worksheet.append(template.headers)
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ══════════════════════════════════════════════════════════════════════
# 값 변환
# ══════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (datetime(2026, 6, 19, 13, 30), date(2026, 6, 19)),
        (date(2026, 6, 19), date(2026, 6, 19)),
        ("2026-06-19", date(2026, 6, 19)),
        ("2026/06/19", date(2026, 6, 19)),
        ("2026.6.19", date(2026, 6, 19)),
        ("2026년 6월 19일", date(2026, 6, 19)),
        ("26-06-19", date(2026, 6, 19)),
        ("2026-06-19 00:00:00", date(2026, 6, 19)),
        (46192, date(2026, 6, 19)),  # 엑셀 일련번호
        (None, None),
        ("", None),
    ],
)
def test_date_conversion_accepts_what_excel_actually_produces(raw, expected):
    assert E._to_date(raw) == expected


@pytest.mark.parametrize("raw", ["어제", "2026-13-01", "2026-02-30", 5])
def test_unreadable_date_reports_instead_of_silently_becoming_none(raw):
    with pytest.raises(E.CellError):
        E._to_date(raw)


def test_money_does_not_pass_through_float():
    """`Decimal(0.1)` 은 0.1000000000000000055... 다. str() 을 거쳐야 한다."""
    assert E._to_decimal(0.1) == Decimal("0.1")
    assert E._to_decimal(520000.55) == Decimal("520000.55")


def test_money_accepts_korean_formatting():
    assert E._to_decimal("1,234,000원") == Decimal("1234000")
    assert E._to_decimal("20%") == Decimal("20")


def test_int_rejects_a_fractional_number():
    assert E._to_int(3.0) == 3
    with pytest.raises(E.CellError):
        E._to_int(3.5)


@pytest.mark.parametrize(
    ("raw", "expected"),
    [("Y", True), ("y", True), ("확정", True), (True, True),
     ("N", False), ("", False), (None, False)],
)
def test_bool_conversion(raw, expected):
    assert E._to_bool(raw) is expected


def test_unknown_bool_value_is_reported():
    with pytest.raises(E.CellError):
        E._to_bool("아마도")


# ══════════════════════════════════════════════════════════════════════
# 양식
# ══════════════════════════════════════════════════════════════════════


def test_every_kind_has_a_template():
    assert set(E.TEMPLATES) == set(E.Kind)


def test_generated_template_round_trips_to_zero_rows():
    """예시 행은 데이터가 아니다. 지우지 않고 올리는 일이 잦다."""
    content = E.generate_template(E.Kind.SALES_ORDER).getvalue()
    result = E.parse(E.Kind.SALES_ORDER, content)
    assert result.rows == []
    assert result.blank_rows == 1
    assert result.ok


def test_full_template_contains_a_sheet_per_kind():
    workbook = load_workbook(E.generate_template())
    assert set(workbook.sheetnames) == {t.sheet_name for t in E.TEMPLATES.values()}


def test_unknown_template_is_rejected():
    with pytest.raises(ValueError, match="지원하지 않는 양식"):
        E.get_template("존재하지않음")


# ══════════════════════════════════════════════════════════════════════
# 파싱
# ══════════════════════════════════════════════════════════════════════


def test_rows_are_read_by_header_name_not_position():
    """사용자가 열 순서를 바꾸거나 메모 열을 끼워 넣는 일이 흔하다."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "판매"
    worksheet.append(["메모", "출고일", "제품명", "채널", "수량"])
    worksheet.append(["아무거나", "2026-06-19", "드리미 H12 Pro", "쿠팡", 5])
    buffer = BytesIO()
    workbook.save(buffer)

    result = E.parse(E.Kind.SALES_ORDER, buffer.getvalue())
    assert result.ok
    assert result.rows[0]["qty"] == 5
    assert result.rows[0]["ship_date"] == date(2026, 6, 19)


def test_missing_required_column_fails_the_whole_file():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "판매"
    worksheet.append(["카테고리", "채널"])  # 제품명 · 수량 · 출고일 없음
    worksheet.append(["가전", "쿠팡"])
    buffer = BytesIO()
    workbook.save(buffer)

    with pytest.raises(ValueError, match="필수 열이 없습니다"):
        E.parse(E.Kind.SALES_ORDER, buffer.getvalue())


def test_bad_row_is_reported_with_its_position_and_others_survive():
    content = build_sheet(
        E.Kind.SALES_ORDER,
        [
            ["가전", "쿠팡", "드리미 H12 Pro", 3, "2026-06-19", None, None, None, None],
            ["가전", "쿠팡", "불량행", "셋", "2026-06-19", None, None, None, None],
            ["가전", "네이버", "드리미 L10", 7, "2026-06-20", None, None, None, None],
        ],
    )
    result = E.parse(E.Kind.SALES_ORDER, content)

    assert len(result.rows) == 2
    assert len(result.errors) == 1
    error = result.errors[0]
    assert error.row_no == 3  # 헤더가 1행이므로 두 번째 데이터 행은 3행
    assert error.column == "수량"


def test_missing_required_value_is_reported():
    content = build_sheet(
        E.Kind.SALES_ORDER,
        [["가전", "쿠팡", None, 3, "2026-06-19", None, None, None, None]],
    )
    result = E.parse(E.Kind.SALES_ORDER, content)
    assert result.rows == []
    assert result.errors[0].column == "제품명"
    assert "필수" in result.errors[0].message


def test_blank_rows_are_counted_not_treated_as_errors():
    content = build_sheet(
        E.Kind.SALES_ORDER,
        [
            ["가전", "쿠팡", "드리미 H12 Pro", 3, "2026-06-19", None, None, None, None],
            [None, None, None, None, None, None, None, None, None],
            ["", "", "", "", "", "", "", "", ""],
        ],
    )
    result = E.parse(E.Kind.SALES_ORDER, content)
    assert len(result.rows) == 1
    assert result.blank_rows == 2
    assert result.ok


def test_corrupt_file_raises_a_readable_error():
    with pytest.raises(ValueError, match="엑셀 파일을 열 수 없습니다"):
        E.parse(E.Kind.SALES_ORDER, b"this is not a workbook")


# ══════════════════════════════════════════════════════════════════════
# 적재 — 판매
# ══════════════════════════════════════════════════════════════════════


@pytest.fixture
def world(db: Session):
    brand = make_brand(db)
    warehouse = make_warehouse(db)
    channel = make_channel(db, name="쿠팡", channel_group="온라인", warehouse=warehouse)
    product = make_product(db, brand, product_code="H12", name="H12 Pro")
    make_alias(db, product, "드리미 H12 Pro", "H12 시리즈")
    return {"brand": brand, "warehouse": warehouse, "channel": channel, "product": product}


def test_sales_upload_maps_known_names_and_derives_weekly_metrics(
    db: Session, world: dict
):
    content = build_sheet(
        E.Kind.SALES_ORDER,
        [["가전", "쿠팡", "드리미 H12 Pro", 5, "2026-06-19", None, "ORD-1", None, "50000"]],
    )
    result = E.import_sales_orders(
        db, E.parse(E.Kind.SALES_ORDER, content), brand_id=world["brand"].id
    )

    assert result.created == 1
    assert result.unresolved == []

    order = db.scalar(select(SalesOrder))
    assert order.product_id == world["product"].id
    assert order.channel_id == world["channel"].id
    assert order.iso_year == 2026
    assert order.iso_week == 25

    # 업로드 한 번으로 주차 집계까지 만들어진다 — 별도 동기화 작업이 없다.
    assert result.rebuild is not None
    metric = db.scalar(select(WeeklyMetric))
    assert metric.sales_qty == 5


def test_unmapped_names_are_stored_and_surfaced(db: Session, world: dict):
    """조용히 버리지 않는다. 나중에 별칭을 추가하면 다시 해석된다."""
    content = build_sheet(
        E.Kind.SALES_ORDER,
        [["가전", "낯선채널", "낯선제품", 5, "2026-06-19", None, None, None, "1000"]],
    )
    result = E.import_sales_orders(
        db, E.parse(E.Kind.SALES_ORDER, content), brand_id=world["brand"].id
    )

    assert result.created == 1
    assert result.unresolved == ["제품: 낯선제품", "채널: 낯선채널"]

    order = db.scalar(select(SalesOrder))
    assert order.product_id is None
    assert order.source_product_name == "낯선제품"


def test_amount_is_computed_from_unit_price_when_missing(db: Session, world: dict):
    content = build_sheet(
        E.Kind.SALES_ORDER,
        [["가전", "쿠팡", "드리미 H12 Pro", 3, "2026-06-19", None, None, "1000.50", None]],
    )
    E.import_sales_orders(
        db, E.parse(E.Kind.SALES_ORDER, content), brand_id=world["brand"].id
    )

    order = db.scalar(select(SalesOrder))
    assert order.amount == Decimal("3001.50")


def test_upload_rebuilds_only_the_uploaded_weeks(db: Session, world: dict):
    content = build_sheet(
        E.Kind.SALES_ORDER,
        [
            ["가전", "쿠팡", "드리미 H12 Pro", 5, "2026-06-19", None, None, None, "1000"],
            ["가전", "쿠팡", "드리미 H12 Pro", 5, "2026-06-22", None, None, None, "1000"],
        ],
    )
    result = E.import_sales_orders(
        db, E.parse(E.Kind.SALES_ORDER, content), brand_id=world["brand"].id
    )
    # 6/19 는 25주차, 6/22 는 26주차다.
    assert result.rebuild.weeks_covered == 2


def test_large_upload_is_written_in_chunks(db: Session, world: dict):
    rows = [
        ["가전", "쿠팡", "드리미 H12 Pro", 1, "2026-06-19", None, f"ORD-{n}", None, "100"]
        for n in range(E.CHUNK_SIZE + 25)
    ]
    result = E.import_sales_orders(
        db, E.parse(E.Kind.SALES_ORDER, build_sheet(E.Kind.SALES_ORDER, rows)),
        brand_id=world["brand"].id,
    )
    assert result.created == E.CHUNK_SIZE + 25
    assert db.query(SalesOrder).count() == E.CHUNK_SIZE + 25


# ══════════════════════════════════════════════════════════════════════
# 적재 — 재고 · 입고 · 행사 · 기준정보
# ══════════════════════════════════════════════════════════════════════


def test_snapshot_upload_updates_the_same_lot_instead_of_duplicating(
    db: Session, world: dict
):
    """같은 (일자, 창고, 품목, 기한)은 한 건이다. 스냅샷은 사실의 기록이다."""
    rows = [["2026-06-19", "용인 메인", "H12", "2028-05-15", 12000]]
    E.import_inventory_snapshots(db, E.parse(E.Kind.INVENTORY_SNAPSHOT, build_sheet(E.Kind.INVENTORY_SNAPSHOT, rows)))

    rows = [["2026-06-19", "용인 메인", "H12", "2028-05-15", 9000]]
    result = E.import_inventory_snapshots(
        db, E.parse(E.Kind.INVENTORY_SNAPSHOT, build_sheet(E.Kind.INVENTORY_SNAPSHOT, rows))
    )

    assert result.updated == 1
    assert db.query(InventorySnapshot).count() == 1
    assert db.scalar(select(InventorySnapshot)).qty == 9000


def test_snapshot_with_unknown_warehouse_is_reported(db: Session, world: dict):
    rows = [["2026-06-19", "없는창고", "H12", None, 100]]
    result = E.import_inventory_snapshots(
        db, E.parse(E.Kind.INVENTORY_SNAPSHOT, build_sheet(E.Kind.INVENTORY_SNAPSHOT, rows))
    )
    assert result.created == 0
    assert "없는창고" in result.errors[0].message


def test_inbound_upload_updates_by_invoice_number(db: Session, world: dict):
    def sheet(status: str):
        return build_sheet(
            E.Kind.INBOUND,
            [[
                "INV-1", "BL-1", "PC-1", "PRD-1", "H12", "용인 메인",
                "2026-03-01", "2026-04-15", "2026-04-10", "2026-02-15", "2028-02-15",
                300, 3600, "8.5", "30600", "1385.5", "42396300", status,
            ]],
        )

    E.import_inbounds(db, E.parse(E.Kind.INBOUND, sheet("해상운송중")))
    result = E.import_inbounds(db, E.parse(E.Kind.INBOUND, sheet("한국도착")))

    assert result.updated == 1
    assert db.query(Inbound).count() == 1
    assert db.scalar(select(Inbound)).status == InboundStatus.ARRIVED_KR


def test_inbound_with_unknown_status_is_reported(db: Session, world: dict):
    content = build_sheet(
        E.Kind.INBOUND,
        [["INV-2", None, None, None, "H12", None, None, None, None, None, None,
          None, 10, None, None, None, None, "우주여행중"]],
    )
    result = E.import_inbounds(db, E.parse(E.Kind.INBOUND, content))
    assert result.created == 0
    assert "우주여행중" in result.errors[0].message


def test_promotion_upload_rejects_reversed_dates(db: Session, world: dict):
    content = build_sheet(
        E.Kind.PROMOTION,
        [["2026-06-21", "2026-06-15", "쿠팡", "메인", "여름 특가", "드리미 H12 Pro",
          None, None, None, None, None, None, "Y", "Y"]],
    )
    result = E.import_promotions(
        db, E.parse(E.Kind.PROMOTION, content), brand_id=world["brand"].id
    )
    assert result.created == 0
    assert "시작일보다" in result.errors[0].message


def test_promotion_upload_resolves_mappings(db: Session, world: dict):
    content = build_sheet(
        E.Kind.PROMOTION,
        [["2026-06-15", "2026-06-21", "쿠팡", "메인 배너", "여름 특가", "드리미 H12 Pro",
          "399000", "319200", "20", None, "필터 2종", None, "Y", "Y"]],
    )
    result = E.import_promotions(
        db, E.parse(E.Kind.PROMOTION, content), brand_id=world["brand"].id
    )

    assert result.created == 1
    assert result.unresolved == []
    promotion = db.scalar(select(Promotion))
    assert promotion.product_id == world["product"].id
    assert promotion.channel_id == world["channel"].id
    assert promotion.price == Decimal("319200.00")
    assert promotion.is_confirmed is True


def test_master_product_upload_creates_then_updates(db: Session, world: dict):
    def sheet(name: str, price: str):
        return build_sheet(
            E.Kind.PRODUCT, [["X30", name, "드리미", 1, "USD", price]]
        )

    first = E.import_master(db, E.parse(E.Kind.PRODUCT, sheet("X30 로봇", "738.5000")))
    assert first.created == 1

    second = E.import_master(db, E.parse(E.Kind.PRODUCT, sheet("X30 로봇청소기", "800.0000")))
    assert second.updated == 1

    from app.models.master import Product

    product = db.scalar(select(Product).where(Product.product_code == "X30"))
    assert product.name == "X30 로봇청소기"
    assert product.purchase_price == Decimal("800.0000")


def test_master_upload_with_unknown_brand_is_reported(db: Session, world: dict):
    content = build_sheet(E.Kind.PRODUCT, [["Z1", "미지품목", "없는브랜드", 1, "USD", "1"]])
    result = E.import_master(db, E.parse(E.Kind.PRODUCT, content))
    assert result.created == 0
    assert "없는브랜드" in result.errors[0].message


def test_channel_upload_links_the_warehouse(db: Session, world: dict):
    content = build_sheet(E.Kind.CHANNEL, [["네이버", "온라인", "용인 메인", "Y"]])
    result = E.import_master(db, E.parse(E.Kind.CHANNEL, content))

    assert result.created == 1
    from app.models.master import Channel

    channel = db.scalar(select(Channel).where(Channel.name == "네이버"))
    assert channel.warehouse_id == world["warehouse"].id
    assert channel.is_major is True


def test_import_master_rejects_a_non_master_kind(db: Session):
    with pytest.raises(ValueError, match="기준 정보가 아닙니다"):
        E.import_master(db, E.ParseResult(kind=E.Kind.SALES_ORDER))
