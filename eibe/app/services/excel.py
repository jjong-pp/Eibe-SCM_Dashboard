"""
엑셀 입출력 — 양식 생성 · 업로드 파싱 · 적재.

구 `app/core/excel_parser.py` 를 대체한다. 달라진 점:

1. **pandas 를 걷어냈다.** 엑셀 한 장을 행 단위로 읽는 데 DataFrame 이 필요하지
   않다. pandas 는 빈 칸을 `NaN`(float) 으로 만들기 때문에 구 코드에는
   `pd.notna(...)` 와 `str(row.iloc[i])` 가 줄줄이 붙어 있었고, 그 과정에서
   날짜와 금액이 문자열로 뭉개졌다. openpyxl 은 셀 타입을 그대로 준다.

2. **금액을 float 으로 거치지 않는다.** openpyxl 은 숫자 셀을 float 으로 주는데
   `Decimal(0.1)` 은 `0.1000000000000000055...` 다. 반드시 `str()` 을 경유한다.
   P-02 와 같은 문제를 입력 단계에서 되풀이하지 않기 위한 것이다.

3. **행 단위로 실패한다.** 한 줄이 잘못됐다고 업로드 전체를 되돌리지 않는다.
   대신 몇 번째 줄의 어느 칸이 왜 틀렸는지 모아 돌려준다 — 구 코드는
   `except Exception: continue` 로 조용히 버렸다.

4. **적재 후 해당 주차만 재집계한다.** 3년치를 매번 훑지 않는다.
"""

from __future__ import annotations

import logging
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from enum import StrEnum
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.dates import iso_week_of
from app.models.enums import InboundStatus, WarehouseType
from app.models.master import Brand, Channel, Product, ProductAlias, Warehouse
from app.models.sales import Promotion, SalesOrder
from app.models.scm import Inbound, InventorySnapshot
from app.services import derive

logger = logging.getLogger(__name__)

#: 한 번에 flush 할 행 수. 수만 행짜리 업로드에서 메모리가 터지지 않게 한다.
CHUNK_SIZE = 500

#: 엑셀 날짜 일련번호의 기준일. 1900년 윤년 버그 때문에 1899-12-30 이다.
_EXCEL_EPOCH = date(1899, 12, 30)

#: 일련번호로 볼 최소값 (20000 ≈ 1954년). 그 아래 숫자는 날짜가 아니라고 본다.
_EXCEL_SERIAL_MIN = 20000
_EXCEL_SERIAL_MAX = 100000


# ══════════════════════════════════════════════════════════════════════
# 양식 정의
# ══════════════════════════════════════════════════════════════════════


class Kind(StrEnum):
    """업로드·양식 종류."""

    PRODUCT = "product"
    WAREHOUSE = "warehouse"
    CHANNEL = "channel"
    INBOUND = "inbound"
    INVENTORY_SNAPSHOT = "inventory_snapshot"
    SALES_ORDER = "sales_order"
    PROMOTION = "promotion"


class FieldKind(StrEnum):
    TEXT = "text"
    INT = "int"
    DECIMAL = "decimal"
    DATE = "date"
    BOOL = "bool"


@dataclass(frozen=True, slots=True)
class Column:
    header: str
    field: str
    kind: FieldKind = FieldKind.TEXT
    required: bool = False
    example: Any = ""


@dataclass(frozen=True, slots=True)
class Template:
    kind: Kind
    sheet_name: str
    columns: tuple[Column, ...]

    @property
    def filename(self) -> str:
        return f"{self.sheet_name}_양식.xlsx"

    @property
    def headers(self) -> list[str]:
        return [column.header for column in self.columns]


def _enum_hint(enum_cls: type[StrEnum]) -> str:
    return "/".join(member.value for member in enum_cls)


TEMPLATES: dict[Kind, Template] = {
    Kind.PRODUCT: Template(
        kind=Kind.PRODUCT,
        sheet_name="품목",
        columns=(
            Column("품목코드", "product_code", required=True, example="H12PRO"),
            Column("품목명", "name", required=True, example="드리미 H12 Pro 무선청소기"),
            Column("브랜드", "brand_name", required=True, example="드리미"),
            Column("카툰당입수량", "pack_qty_per_tu", FieldKind.INT, example=1),
            Column("통화", "currency", example="USD"),
            Column("매입가", "purchase_price", FieldKind.DECIMAL, example="289.0000"),
        ),
    ),
    Kind.WAREHOUSE: Template(
        kind=Kind.WAREHOUSE,
        sheet_name="창고",
        columns=(
            Column("창고명", "name", required=True, example="용인 메인"),
            Column(
                f"창고타입({_enum_hint(WarehouseType)})", "type", example="HUB"
            ),
            Column("허용잔여기한일수", "allowed_expiry_days", FieldKind.INT, example=90),
            Column("기본이관MOQ", "default_transfer_moq", FieldKind.INT, example=0),
        ),
    ),
    Kind.CHANNEL: Template(
        kind=Kind.CHANNEL,
        sheet_name="채널",
        columns=(
            Column("채널명", "name", required=True, example="쿠팡"),
            Column("채널그룹", "channel_group", example="온라인"),
            Column("연결창고", "warehouse_name", example="쿠팡 FFC"),
            Column("주요채널", "is_major", FieldKind.BOOL, example="Y"),
        ),
    ),
    Kind.INBOUND: Template(
        kind=Kind.INBOUND,
        sheet_name="입고",
        columns=(
            Column("입고번호", "invoice_no", example="INV-2026-0001"),
            Column("BL번호", "bl_no", example="BL0010001"),
            Column("구매코드", "purchase_code", example="PC-H12PRO-1"),
            Column("생산코드", "production_code", example="PRD-H12PRO-1"),
            Column("품목코드", "product_code", required=True, example="H12PRO"),
            Column("도착창고", "warehouse_name", example="용인 메인"),
            Column("선적일", "shipping_date", FieldKind.DATE, example="2026-03-01"),
            Column("한국도착일", "korea_arrival_date", FieldKind.DATE, example="2026-04-15"),
            Column("ETA", "eta", FieldKind.DATE, example="2026-04-10"),
            Column("제조일자", "manufacture_date", FieldKind.DATE, example="2026-02-15"),
            Column("기한", "expiry_date", FieldKind.DATE, example="2028-02-15"),
            Column("카툰수", "carton_qty", FieldKind.INT, example=300),
            Column("수량", "unit_qty", FieldKind.INT, required=True, example=3600),
            Column("단가", "unit_price", FieldKind.DECIMAL, example="8.5000"),
            Column("총액", "total_price", FieldKind.DECIMAL, example="30600.00"),
            Column("결제환율", "exchange_rate", FieldKind.DECIMAL, example="1385.5000"),
            Column("결제금액(원화)", "payment_amount_krw", FieldKind.DECIMAL, example="42396300.00"),
            Column(f"상태({_enum_hint(InboundStatus)})", "status", example="해상운송중"),
        ),
    ),
    Kind.INVENTORY_SNAPSHOT: Template(
        kind=Kind.INVENTORY_SNAPSHOT,
        sheet_name="현재고",
        columns=(
            Column("스냅샷일자", "snapshot_date", FieldKind.DATE, required=True, example="2026-06-19"),
            Column("창고명", "warehouse_name", required=True, example="용인 메인"),
            Column("품목코드", "product_code", required=True, example="H12PRO"),
            Column("기한", "expiry_date", FieldKind.DATE, example="2028-05-15"),
            Column("수량", "qty", FieldKind.INT, required=True, example=12000),
        ),
    ),
    Kind.SALES_ORDER: Template(
        kind=Kind.SALES_ORDER,
        sheet_name="판매",
        columns=(
            Column("카테고리", "category", example="가전"),
            Column("채널", "source_channel_name", required=True, example="쿠팡"),
            Column("제품명", "source_product_name", required=True, example="드리미 H12 Pro"),
            Column("수량", "qty", FieldKind.INT, required=True, example=3),
            Column("출고일", "ship_date", FieldKind.DATE, required=True, example="2026-06-19"),
            Column("주문일자", "order_date", FieldKind.DATE, example="2026-06-18"),
            Column("주문번호", "order_no", example="ORD-20260619-0001"),
            Column("판매단가", "unit_price", FieldKind.DECIMAL, example="520000.00"),
            Column("판매금액", "amount", FieldKind.DECIMAL, example="1560000.00"),
        ),
    ),
    Kind.PROMOTION: Template(
        kind=Kind.PROMOTION,
        sheet_name="행사",
        columns=(
            Column("시작일", "start_date", FieldKind.DATE, required=True, example="2026-06-15"),
            Column("종료일", "end_date", FieldKind.DATE, required=True, example="2026-06-21"),
            Column("채널", "source_channel_name", required=True, example="쿠팡"),
            Column("구좌명", "slot_name", example="메인 배너"),
            Column("행사명", "event_name", required=True, example="여름 특가"),
            Column("제품명", "source_product_name", example="드리미 H12 Pro"),
            Column("정상가", "list_price", FieldKind.DECIMAL, example="399000.00"),
            Column("행사가", "price", FieldKind.DECIMAL, example="319200.00"),
            Column("할인율", "discount_rate", FieldKind.DECIMAL, example="20.000"),
            Column("적립금", "reward_points", FieldKind.DECIMAL, example="0.00"),
            Column("사은품", "gift", example="사은품 필터 2종"),
            Column("비고", "note", example=""),
            Column("마케팅여부", "is_marketing", FieldKind.BOOL, example="Y"),
            Column("확정여부", "is_confirmed", FieldKind.BOOL, example="Y"),
        ),
    ),
}


# ══════════════════════════════════════════════════════════════════════
# 양식 생성
# ══════════════════════════════════════════════════════════════════════

_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
# 긍정 색 #29AD3A — UI 원칙과 같은 값을 쓴다.
_HEADER_FILL = PatternFill("solid", start_color="29AD3A", end_color="29AD3A")
_EXAMPLE_FONT = Font(italic=True, color="999999", size=10)
_EXAMPLE_FILL = PatternFill("solid", start_color="F5F5F5", end_color="F5F5F5")
_THIN = Side(style="thin", color="D4D4D4")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _write_sheet(worksheet, template: Template) -> None:
    for index, column in enumerate(template.columns, start=1):
        cell = worksheet.cell(row=1, column=index, value=column.header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
        # 한글은 폭을 넓게 잡아야 잘리지 않는다.
        worksheet.column_dimensions[get_column_letter(index)].width = max(
            len(column.header) * 2.0, 14
        )

        example = worksheet.cell(row=2, column=index, value=column.example)
        example.font = _EXAMPLE_FONT
        example.fill = _EXAMPLE_FILL
        example.alignment = Alignment(horizontal="center")
        example.border = _BORDER

    worksheet.freeze_panes = "A3"


def generate_template(kind: Kind | str | None = None) -> BytesIO:
    """엑셀 양식을 만든다. kind 를 생략하면 모든 시트를 한 파일에 담는다."""
    if kind is None:
        templates = list(TEMPLATES.values())
    else:
        templates = [get_template(kind)]

    workbook = Workbook()
    workbook.remove(workbook.active)
    for template in templates:
        _write_sheet(workbook.create_sheet(template.sheet_name), template)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def get_template(kind: Kind | str) -> Template:
    try:
        return TEMPLATES[Kind(kind)]
    except ValueError as exc:
        raise ValueError(f"지원하지 않는 양식입니다: {kind}") from exc


def template_filename(kind: Kind | str | None = None) -> str:
    if kind is None:
        return "EIBE_데이터입력양식_전체.xlsx"
    return get_template(kind).filename


# ══════════════════════════════════════════════════════════════════════
# 값 변환
# ══════════════════════════════════════════════════════════════════════


class CellError(ValueError):
    """한 칸을 해석하지 못했다. 메시지가 그대로 사용자에게 간다."""


def _to_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _to_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        raise CellError("숫자가 필요합니다.")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not value.is_integer():
            raise CellError(f"정수가 필요합니다: {value}")
        return int(value)
    text = str(value).strip().replace(",", "").replace("개", "")
    try:
        return int(Decimal(text))
    except (InvalidOperation, ValueError) as exc:
        raise CellError(f"숫자로 읽을 수 없습니다: {value!r}") from exc


def _to_decimal(value: Any) -> Decimal | None:
    """금액·비율. **float 을 직접 Decimal 에 넣지 않는다.**

    `Decimal(0.1)` 은 `0.1000000000000000055511151231257827` 이다. openpyxl 이
    숫자 셀을 float 으로 주므로 반드시 `str()` 을 거쳐야 한다.
    """
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        raise CellError("금액이 필요합니다.")
    if isinstance(value, (int, float)):
        return Decimal(str(value))

    text = str(value).strip().replace(",", "").replace("원", "").replace("%", "")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise CellError(f"금액으로 읽을 수 없습니다: {value!r}") from exc


def _to_date(value: Any) -> date | None:
    """날짜. 엑셀 일련번호 · datetime · 한국식 표기를 모두 받는다."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    # 서식이 '일반'인 날짜 칸은 숫자로 들어온다.
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        serial = int(value)
        if _EXCEL_SERIAL_MIN <= serial <= _EXCEL_SERIAL_MAX:
            return _EXCEL_EPOCH + timedelta(days=serial)
        raise CellError(f"날짜로 읽을 수 없습니다: {value!r}")

    text = str(value).strip()

    # '2026-06-19 00:00:00' 처럼 시간이 붙어 오는 경우 앞부분만 쓴다.
    # 공백으로 무조건 자르면 안 된다 — '2026년 6월 19일'도 공백을 포함하므로
    # 시간이 실제로 있을 때(':' 가 있을 때)만 자른다.
    if ":" in text:
        text = text.replace("T", " ").split(" ")[0]

    # '2026년 6월 19일' → '2026-6-19'
    text = text.replace("년", "-").replace("월", "-").replace("일", "")
    for separator in ("/", "."):
        text = text.replace(separator, "-")
    text = "".join(text.split()).strip("-")

    parts = [part for part in text.split("-") if part]
    if len(parts) == 3:
        try:
            year, month, day = (int(part) for part in parts)
        except ValueError as exc:
            raise CellError(f"날짜로 읽을 수 없습니다: {value!r}") from exc
        if year < 100:  # '26-06-19'
            year += 2000
        try:
            return date(year, month, day)
        except ValueError as exc:
            raise CellError(f"존재하지 않는 날짜입니다: {value!r}") from exc

    raise CellError(f"날짜로 읽을 수 없습니다: {value!r}")


_TRUE_VALUES = {"y", "yes", "true", "1", "o", "확정", "예", "사용"}
_FALSE_VALUES = {"n", "no", "false", "0", "x", "미확정", "아니오", "미사용", ""}


def _to_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    text = str(value).strip().lower()
    if text in _TRUE_VALUES:
        return True
    if text in _FALSE_VALUES:
        return False
    raise CellError(f"Y 또는 N 이어야 합니다: {value!r}")


_CONVERTERS = {
    FieldKind.TEXT: _to_text,
    FieldKind.INT: _to_int,
    FieldKind.DECIMAL: _to_decimal,
    FieldKind.DATE: _to_date,
    FieldKind.BOOL: _to_bool,
}


# ══════════════════════════════════════════════════════════════════════
# 파싱
# ══════════════════════════════════════════════════════════════════════


@dataclass(frozen=True, slots=True)
class RowError:
    """어느 줄의 어느 칸이 왜 틀렸는가. 그대로 화면에 띄운다."""

    row_no: int
    column: str
    message: str

    def __str__(self) -> str:
        return f"{self.row_no}행 [{self.column}] {self.message}"


@dataclass(frozen=True, slots=True)
class ParseResult:
    kind: Kind
    rows: list[dict[str, Any]] = field(default_factory=list)
    errors: list[RowError] = field(default_factory=list)
    #: 빈 줄이라 건너뛴 수. 엑셀 하단의 서식만 남은 행이 여기 잡힌다.
    blank_rows: int = 0

    @property
    def ok(self) -> bool:
        return not self.errors


def _header_index(worksheet, template: Template) -> dict[str, int]:
    """헤더 이름 → 열 번호.

    열 순서가 바뀌어도 이름으로 찾는다. 사용자가 열을 옮기거나 중간에
    메모 열을 끼워 넣는 일이 흔하다.
    """
    headers: dict[str, int] = {}
    for index, cell in enumerate(next(worksheet.iter_rows(max_row=1)), start=1):
        name = _to_text(cell.value)
        if name:
            headers[name] = index

    missing = [
        column.header
        for column in template.columns
        if column.required and column.header not in headers
    ]
    if missing:
        raise ValueError(f"필수 열이 없습니다: {', '.join(missing)}")
    return headers


def _is_example_row(values: dict[str, Any], template: Template) -> bool:
    """양식에 딸린 회색 예시 행인가.

    사용자가 예시를 지우지 않고 그대로 올리는 일이 잦다.
    """
    filled = [
        (column, values.get(column.field))
        for column in template.columns
        if column.example not in ("", None)
    ]
    if not filled:
        return False
    return all(
        value is not None and str(value).strip() == str(column.example).strip()
        for column, value in filled
    )


def parse(kind: Kind | str, content: bytes) -> ParseResult:
    """업로드된 엑셀을 행 dict 목록으로 읽는다.

    한 줄이 잘못돼도 나머지는 읽는다. 오류는 `ParseResult.errors` 에 모인다.
    """
    template = get_template(kind)

    try:
        # read_only 로 큰 파일도 메모리에 통째로 올리지 않는다.
        # data_only=True 는 수식 대신 계산된 값을 읽는다.
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 은 다양한 예외를 던진다
        raise ValueError(f"엑셀 파일을 열 수 없습니다: {exc}") from exc

    try:
        worksheet = (
            workbook[template.sheet_name]
            if template.sheet_name in workbook.sheetnames
            else workbook[workbook.sheetnames[0]]
        )
        if worksheet.max_row is None or worksheet.max_row < 1:
            return ParseResult(kind=template.kind)

        headers = _header_index(worksheet, template)

        rows: list[dict[str, Any]] = []
        errors: list[RowError] = []
        blank = 0

        for row_no, cells in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            if all(cell is None or str(cell).strip() == "" for cell in cells):
                blank += 1
                continue

            record: dict[str, Any] = {}
            row_errors: list[RowError] = []

            for column in template.columns:
                position = headers.get(column.header)
                raw = cells[position - 1] if position and position <= len(cells) else None

                try:
                    value = _CONVERTERS[column.kind](raw)
                except CellError as exc:
                    row_errors.append(RowError(row_no, column.header, str(exc)))
                    continue

                if value is None and column.required:
                    row_errors.append(RowError(row_no, column.header, "필수 값입니다."))
                    continue

                record[column.field] = value

            if _is_example_row(record, template):
                blank += 1
                continue

            if row_errors:
                errors.extend(row_errors)
                continue

            rows.append(record)

        return ParseResult(kind=template.kind, rows=rows, errors=errors, blank_rows=blank)
    finally:
        workbook.close()


# ══════════════════════════════════════════════════════════════════════
# 적재
# ══════════════════════════════════════════════════════════════════════


@dataclass(frozen=True, slots=True)
class ImportResult:
    kind: Kind
    created: int = 0
    updated: int = 0
    errors: list[RowError] = field(default_factory=list)
    #: 매핑되지 않은 원본 값. 적재는 되지만 집계에서 빠지므로 드러내야 한다.
    unresolved: list[str] = field(default_factory=list)
    rebuild: derive.RebuildResult | None = None

    def __str__(self) -> str:
        summary = f"{self.kind.value}: 신규 {self.created}건, 갱신 {self.updated}건"
        if self.errors:
            summary += f", 오류 {len(self.errors)}건"
        if self.unresolved:
            summary += f", 미매핑 {len(self.unresolved)}종"
        return summary


def _chunked(rows: list[dict[str, Any]]) -> Iterator[list[dict[str, Any]]]:
    for start in range(0, len(rows), CHUNK_SIZE):
        yield rows[start : start + CHUNK_SIZE]


def import_sales_orders(
    db: Session, result: ParseResult, *, brand_id: int, rebuild: bool = True
) -> ImportResult:
    """판매 원장 적재.

    매핑되지 않은 제품·채널도 **원본 문자열 그대로 저장한다.** 나중에 별칭을
    추가하면 `derive.resolve_sales_mappings()` 로 다시 해석되므로 재업로드가
    필요 없다. 대신 미매핑 목록을 돌려주어 조용히 사라지지 않게 한다.
    """
    alias_map = {
        alias.source_name: alias.product_id for alias in db.scalars(select(ProductAlias))
    }
    channel_map = {channel.name: channel.id for channel in db.scalars(select(Channel))}

    created = 0
    unresolved: set[str] = set()
    touched: list[date] = []

    for chunk in _chunked(result.rows):
        for record in chunk:
            ship_date = record["ship_date"]
            week = iso_week_of(ship_date)
            source_product = record["source_product_name"]
            source_channel = record["source_channel_name"]

            product_id = alias_map.get(source_product)
            channel_id = channel_map.get(source_channel)
            if product_id is None:
                unresolved.add(f"제품: {source_product}")
            if channel_id is None:
                unresolved.add(f"채널: {source_channel}")

            qty = record["qty"]
            unit_price = record.get("unit_price")
            amount = record.get("amount")
            if amount is None:
                # 금액이 비어 있으면 단가 × 수량으로 채운다. 둘 다 없으면 0이다.
                amount = unit_price * qty if unit_price is not None else Decimal("0")

            db.add(
                SalesOrder(
                    brand_id=brand_id,
                    order_no=record.get("order_no"),
                    source_product_name=source_product,
                    source_channel_name=source_channel,
                    category=record.get("category"),
                    product_id=product_id,
                    channel_id=channel_id,
                    order_date=record.get("order_date"),
                    ship_date=ship_date,
                    iso_year=week.year,
                    iso_week=week.week,
                    qty=qty,
                    unit_price=unit_price,
                    amount=amount,
                )
            )
            created += 1
            touched.append(ship_date)
        db.flush()

    db.commit()

    rebuilt = None
    if rebuild and touched:
        # 올린 주차만 다시 계산한다. 전체 재계산은 업로드 응답을 느리게 만든다.
        rebuilt = derive.rebuild_weekly_metrics(
            db, since=min(touched), until=max(touched)
        )

    outcome = ImportResult(
        kind=Kind.SALES_ORDER,
        created=created,
        errors=result.errors,
        unresolved=sorted(unresolved),
        rebuild=rebuilt,
    )
    logger.info("판매 원장 업로드 — %s", outcome)
    return outcome


def import_inventory_snapshots(db: Session, result: ParseResult) -> ImportResult:
    """재고 스냅샷 적재.

    같은 (일자, 창고, 품목, 기한) 이 다시 올라오면 수량을 갱신한다. 스냅샷은
    사실의 기록이라 같은 날짜에 두 값이 있으면 안 된다.
    """
    warehouses = {w.name: w.id for w in db.scalars(select(Warehouse))}
    products = {p.product_code: p.id for p in db.scalars(select(Product))}

    created = updated = 0
    errors = list(result.errors)

    for chunk in _chunked(result.rows):
        for index, record in enumerate(chunk):
            warehouse_id = warehouses.get(record["warehouse_name"])
            product_id = products.get(record["product_code"])
            if warehouse_id is None:
                errors.append(
                    RowError(index, "창고명", f"등록되지 않은 창고: {record['warehouse_name']}")
                )
                continue
            if product_id is None:
                errors.append(
                    RowError(index, "품목코드", f"등록되지 않은 품목: {record['product_code']}")
                )
                continue

            existing = db.scalar(
                select(InventorySnapshot).where(
                    InventorySnapshot.snapshot_date == record["snapshot_date"],
                    InventorySnapshot.warehouse_id == warehouse_id,
                    InventorySnapshot.product_id == product_id,
                    InventorySnapshot.expiry_date == record.get("expiry_date"),
                )
            )
            if existing is not None:
                existing.qty = record["qty"]
                updated += 1
            else:
                db.add(
                    InventorySnapshot(
                        snapshot_date=record["snapshot_date"],
                        warehouse_id=warehouse_id,
                        product_id=product_id,
                        expiry_date=record.get("expiry_date"),
                        qty=record["qty"],
                    )
                )
                created += 1
        db.flush()

    db.commit()
    outcome = ImportResult(
        kind=Kind.INVENTORY_SNAPSHOT, created=created, updated=updated, errors=errors
    )
    logger.info("재고 스냅샷 업로드 — %s", outcome)
    return outcome


def import_promotions(db: Session, result: ParseResult, *, brand_id: int) -> ImportResult:
    """행사 적재. 판매와 같은 규칙으로 원본 문자열을 남긴다."""
    alias_map = {
        alias.source_name: alias.product_id for alias in db.scalars(select(ProductAlias))
    }
    channel_map = {channel.name: channel.id for channel in db.scalars(select(Channel))}

    created = 0
    errors = list(result.errors)
    unresolved: set[str] = set()

    for chunk in _chunked(result.rows):
        for index, record in enumerate(chunk):
            if record["end_date"] < record["start_date"]:
                errors.append(RowError(index, "종료일", "시작일보다 빠릅니다."))
                continue

            source_product = record.get("source_product_name")
            source_channel = record["source_channel_name"]
            if source_product and source_product not in alias_map:
                unresolved.add(f"제품: {source_product}")
            if source_channel not in channel_map:
                unresolved.add(f"채널: {source_channel}")

            db.add(
                Promotion(
                    brand_id=brand_id,
                    start_date=record["start_date"],
                    end_date=record["end_date"],
                    source_product_name=source_product,
                    source_channel_name=source_channel,
                    product_id=alias_map.get(source_product) if source_product else None,
                    channel_id=channel_map.get(source_channel),
                    slot_name=record.get("slot_name"),
                    event_name=record["event_name"],
                    list_price=record.get("list_price"),
                    price=record.get("price"),
                    discount_rate=record.get("discount_rate"),
                    reward_points=record.get("reward_points"),
                    gift=record.get("gift"),
                    note=record.get("note"),
                    is_marketing=record.get("is_marketing", False),
                    is_confirmed=record.get("is_confirmed", False),
                )
            )
            created += 1
        db.flush()

    db.commit()
    outcome = ImportResult(
        kind=Kind.PROMOTION, created=created, errors=errors, unresolved=sorted(unresolved)
    )
    logger.info("행사 업로드 — %s", outcome)
    return outcome


def import_inbounds(db: Session, result: ParseResult) -> ImportResult:
    """입고 적재. 입고번호가 같으면 갱신한다 — 진행 상태가 계속 바뀐다."""
    warehouses = {w.name: w.id for w in db.scalars(select(Warehouse))}
    products = {p.product_code: p.id for p in db.scalars(select(Product))}
    statuses = {status.value for status in InboundStatus}

    created = updated = 0
    errors = list(result.errors)

    for chunk in _chunked(result.rows):
        for index, record in enumerate(chunk):
            product_id = products.get(record["product_code"])
            if product_id is None:
                errors.append(
                    RowError(index, "품목코드", f"등록되지 않은 품목: {record['product_code']}")
                )
                continue

            status = record.get("status")
            if status and status not in statuses:
                errors.append(RowError(index, "상태", f"알 수 없는 상태: {status}"))
                continue

            values = {
                "product_id": product_id,
                "arrival_warehouse_id": warehouses.get(record.get("warehouse_name")),
                "bl_no": record.get("bl_no"),
                "purchase_code": record.get("purchase_code"),
                "production_code": record.get("production_code"),
                "shipping_date": record.get("shipping_date"),
                "korea_arrival_date": record.get("korea_arrival_date"),
                "eta": record.get("eta"),
                "manufacture_date": record.get("manufacture_date"),
                "expiry_date": record.get("expiry_date"),
                "carton_qty": record.get("carton_qty"),
                "unit_qty": record.get("unit_qty") or 0,
                "unit_price": record.get("unit_price"),
                "total_price": record.get("total_price"),
                "exchange_rate": record.get("exchange_rate"),
                "payment_amount_krw": record.get("payment_amount_krw"),
                "status": status or InboundStatus.DEPARTED,
            }

            invoice_no = record.get("invoice_no")
            existing = (
                db.scalar(select(Inbound).where(Inbound.invoice_no == invoice_no))
                if invoice_no
                else None
            )
            if existing is not None:
                for attribute, value in values.items():
                    setattr(existing, attribute, value)
                updated += 1
            else:
                db.add(Inbound(invoice_no=invoice_no, **values))
                created += 1
        db.flush()

    db.commit()
    outcome = ImportResult(
        kind=Kind.INBOUND, created=created, updated=updated, errors=errors
    )
    logger.info("입고 업로드 — %s", outcome)
    return outcome


def import_master(db: Session, result: ParseResult) -> ImportResult:
    """기준 정보(품목 · 창고 · 채널) 적재. 업무상 식별자로 갱신한다."""
    created = updated = 0
    errors = list(result.errors)

    if result.kind is Kind.PRODUCT:
        brands = {brand.name: brand.id for brand in db.scalars(select(Brand))}
        for index, record in enumerate(result.rows):
            brand_id = brands.get(record["brand_name"])
            if brand_id is None:
                errors.append(
                    RowError(index, "브랜드", f"등록되지 않은 브랜드: {record['brand_name']}")
                )
                continue
            existing = db.scalar(
                select(Product).where(Product.product_code == record["product_code"])
            )
            values = {
                "name": record["name"],
                "brand_id": brand_id,
                "pack_qty_per_tu": record.get("pack_qty_per_tu") or 1,
                "currency": record.get("currency") or "USD",
                "purchase_price": record.get("purchase_price") or Decimal("0"),
            }
            if existing is not None:
                for attribute, value in values.items():
                    setattr(existing, attribute, value)
                updated += 1
            else:
                db.add(Product(product_code=record["product_code"], **values))
                created += 1

    elif result.kind is Kind.WAREHOUSE:
        for index, record in enumerate(result.rows):
            raw_type = record.get("type") or WarehouseType.ONLINE.value
            try:
                warehouse_type = WarehouseType(raw_type)
            except ValueError:
                errors.append(RowError(index, "창고타입", f"알 수 없는 타입: {raw_type}"))
                continue
            existing = db.scalar(
                select(Warehouse).where(Warehouse.name == record["name"])
            )
            values = {
                "type": warehouse_type,
                "allowed_expiry_days": record.get("allowed_expiry_days") or 0,
                "default_transfer_moq": record.get("default_transfer_moq") or 0,
            }
            if existing is not None:
                for attribute, value in values.items():
                    setattr(existing, attribute, value)
                updated += 1
            else:
                db.add(Warehouse(name=record["name"], **values))
                created += 1

    elif result.kind is Kind.CHANNEL:
        warehouses = {w.name: w.id for w in db.scalars(select(Warehouse))}
        for record in result.rows:
            existing = db.scalar(select(Channel).where(Channel.name == record["name"]))
            values = {
                "channel_group": record.get("channel_group"),
                "warehouse_id": warehouses.get(record.get("warehouse_name")),
                "is_major": record.get("is_major", False),
            }
            if existing is not None:
                for attribute, value in values.items():
                    setattr(existing, attribute, value)
                updated += 1
            else:
                db.add(Channel(name=record["name"], **values))
                created += 1

    else:
        raise ValueError(f"기준 정보가 아닙니다: {result.kind}")

    db.commit()
    outcome = ImportResult(
        kind=result.kind, created=created, updated=updated, errors=errors
    )
    logger.info("기준 정보 업로드 — %s", outcome)
    return outcome
